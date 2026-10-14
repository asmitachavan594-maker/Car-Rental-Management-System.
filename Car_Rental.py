import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os
from datetime import datetime


# ================= EXCEL DATABASE =================

FILE_NAME = "car_rental_data.xlsx"


def create_excel():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()

        ws_cars = wb.active
        ws_cars.title = "Cars"

        ws_cars.append([
            "Car ID",
            "Car Name",
            "Brand",
            "Registration No",
            "Car Type",
            "Price/Day",
            "Availability"
        ])

        ws_rentals = wb.create_sheet("Rentals")

        ws_rentals.append([
            "Rental ID",
            "Customer Name",
            "Contact",
            "Car ID",
            "Rental Date",
            "Return Date",
            "Total Days",
            "Total Amount"
        ])

        wb.save(FILE_NAME)


# ================= LOGIN =================

USERNAME = "admin"
PASSWORD = "1234"


def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == USERNAME and password == PASSWORD:
        login_window.destroy()
        dashboard()
    else:
        messagebox.showerror("Login Failed", "Invalid username or password")


# ================= DASHBOARD =================

def dashboard():
    global root

    root = tk.Tk()
    root.title("Car Rental Management System")
    root.geometry("1000x650")
    root.resizable(False, False)

    title = tk.Label(
        root,
        text="CAR RENTAL MANAGEMENT SYSTEM",
        font=("Arial", 24, "bold")
    )
    title.pack(pady=20)

    menu_frame = tk.Frame(root)
    menu_frame.pack(pady=20)

    tk.Button(
        menu_frame,
        text="Add Car",
        width=18,
        height=2,
        command=add_car
    ).grid(row=0, column=0, padx=10, pady=10)

    tk.Button(
        menu_frame,
        text="View Cars",
        width=18,
        height=2,
        command=view_cars
    ).grid(row=0, column=1, padx=10, pady=10)

    tk.Button(
        menu_frame,
        text="Search Car",
        width=18,
        height=2,
        command=search_car
    ).grid(row=1, column=0, padx=10, pady=10)

    tk.Button(
        menu_frame,
        text="Update Car",
        width=18,
        height=2,
        command=update_car
    ).grid(row=1, column=1, padx=10, pady=10)

    tk.Button(
        menu_frame,
        text="Delete Car",
        width=18,
        height=2,
        command=delete_car
    ).grid(row=2, column=0, padx=10, pady=10)

    tk.Button(
        menu_frame,
        text="Rental",
        width=18,
        height=2,
        command=rental_window
    ).grid(row=2, column=1, padx=10, pady=10)

    tk.Button(
        root,
        text="Logout",
        width=20,
        height=2,
        command=logout
    ).pack(pady=30)

    root.mainloop()


# ================= ADD CAR =================

def add_car():
    window = tk.Toplevel(root)
    window.title("Add Car")
    window.geometry("450x550")

    labels = [
        "Car ID",
        "Car Name",
        "Brand",
        "Registration No",
        "Car Type",
        "Price/Day"
    ]

    entries = []

    for i, label in enumerate(labels):
        tk.Label(
            window,
            text=label,
            font=("Arial", 11)
        ).pack(pady=5)

        entry = tk.Entry(window, width=35)
        entry.pack()

        entries.append(entry)

    def save_car():

        values = [entry.get() for entry in entries]

        if "" in values:
            messagebox.showwarning(
                "Warning",
                "Please fill all fields"
            )
            return

        try:
            price = float(values[5])
        except ValueError:
            messagebox.showerror(
                "Error",
                "Price must be a number"
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb["Cars"]

        ws.append([
            values[0],
            values[1],
            values[2],
            values[3],
            values[4],
            price,
            "Available"
        ])

        wb.save(FILE_NAME)

        messagebox.showinfo(
            "Success",
            "Car added successfully!"
        )

        window.destroy()

    tk.Button(
        window,
        text="Save Car",
        width=20,
        command=save_car
    ).pack(pady=20)


# ================= VIEW CARS =================

def view_cars():
    window = tk.Toplevel(root)
    window.title("View Cars")
    window.geometry("900x450")

    columns = (
        "Car ID",
        "Car Name",
        "Brand",
        "Registration No",
        "Car Type",
        "Price/Day",
        "Availability"
    )

    tree = ttk.Treeview(
        window,
        columns=columns,
        show="headings"
    )

    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=120)

    tree.pack(fill="both", expand=True)

    wb = load_workbook(FILE_NAME)
    ws = wb["Cars"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", "end", values=row)


# ================= SEARCH CAR =================

def search_car():

    window = tk.Toplevel(root)
    window.title("Search Car")
    window.geometry("800x500")

    tk.Label(
        window,
        text="Enter Car ID:"
    ).pack(pady=10)

    search_entry = tk.Entry(window)
    search_entry.pack()

    columns = (
        "Car ID",
        "Car Name",
        "Brand",
        "Registration No",
        "Car Type",
        "Price/Day",
        "Availability"
    )

    tree = ttk.Treeview(
        window,
        columns=columns,
        show="headings"
    )

    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=110)

    tree.pack(fill="both", expand=True, pady=20)

    def search():

        for item in tree.get_children():
            tree.delete(item)

        car_id = search_entry.get()

        wb = load_workbook(FILE_NAME)
        ws = wb["Cars"]

        found = False

        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if str(row[0]) == car_id:
                tree.insert("", "end", values=row)
                found = True

        if not found:
            messagebox.showinfo(
                "Search",
                "Car not found"
            )

    tk.Button(
        window,
        text="Search",
        command=search
    ).pack()


# ================= UPDATE CAR =================

def update_car():

    window = tk.Toplevel(root)
    window.title("Update Car")
    window.geometry("500x600")

    tk.Label(
        window,
        text="Enter Car ID"
    ).pack(pady=10)

    id_entry = tk.Entry(window)
    id_entry.pack()

    fields = [
        "Car Name",
        "Brand",
        "Registration No",
        "Car Type",
        "Price/Day",
        "Availability"
    ]

    entries = {}

    for field in fields:

        tk.Label(
            window,
            text=field
        ).pack(pady=5)

        entry = tk.Entry(window)
        entry.pack()

        entries[field] = entry

    def update():

        car_id = id_entry.get()

        wb = load_workbook(FILE_NAME)
        ws = wb["Cars"]

        found = False

        for row in ws.iter_rows(min_row=2):

            if str(row[0].value) == car_id:

                row[1].value = entries["Car Name"].get()
                row[2].value = entries["Brand"].get()
                row[3].value = entries["Registration No"].get()
                row[4].value = entries["Car Type"].get()

                try:
                    row[5].value = float(
                        entries["Price/Day"].get()
                    )
                except ValueError:
                    messagebox.showerror(
                        "Error",
                        "Price must be a number"
                    )
                    return

                row[6].value = entries["Availability"].get()

                found = True
                break

        if found:

            wb.save(FILE_NAME)

            messagebox.showinfo(
                "Success",
                "Car updated successfully!"
            )

            window.destroy()

        else:
            messagebox.showerror(
                "Error",
                "Car ID not found"
            )

    tk.Button(
        window,
        text="Update",
        width=20,
        command=update
    ).pack(pady=20)


# ================= DELETE CAR =================

def delete_car():

    window = tk.Toplevel(root)
    window.title("Delete Car")
    window.geometry("400x250")

    tk.Label(
        window,
        text="Enter Car ID to delete"
    ).pack(pady=20)

    id_entry = tk.Entry(window)
    id_entry.pack()

    def delete():

        car_id = id_entry.get()

        wb = load_workbook(FILE_NAME)
        ws = wb["Cars"]

        found = False

        for row in range(2, ws.max_row + 1):

            if str(ws.cell(row, 1).value) == car_id:

                confirm = messagebox.askyesno(
                    "Confirm Delete",
                    "Are you sure you want to delete this car?"
                )

                if confirm:
                    ws.delete_rows(row)
                    wb.save(FILE_NAME)

                    messagebox.showinfo(
                        "Success",
                        "Car deleted successfully!"
                    )

                    window.destroy()

                found = True
                break

        if not found:
            messagebox.showerror(
                "Error",
                "Car ID not found"
            )

    tk.Button(
        window,
        text="Delete",
        width=20,
        command=delete
    ).pack(pady=20)


# ================= RENTAL =================

def rental_window():

    window = tk.Toplevel(root)
    window.title("Car Rental")
    window.geometry("500x600")

    fields = [
        "Rental ID",
        "Customer Name",
        "Contact",
        "Car ID",
        "Rental Date",
        "Return Date"
    ]

    entries = {}

    for field in fields:

        tk.Label(
            window,
            text=field
        ).pack(pady=5)

        entry = tk.Entry(window)
        entry.pack()

        entries[field] = entry

    def rent_car():

        try:
            rental_date = datetime.strptime(
                entries["Rental Date"].get(),
                "%d-%m-%Y"
            )

            return_date = datetime.strptime(
                entries["Return Date"].get(),
                "%d-%m-%Y"
            )

            total_days = (
                return_date - rental_date
            ).days

            if total_days <= 0:
                messagebox.showerror(
                    "Error",
                    "Return date must be after rental date"
                )
                return

        except ValueError:
            messagebox.showerror(
                "Error",
                "Use date format DD-MM-YYYY"
            )
            return

        car_id = entries["Car ID"].get()

        wb = load_workbook(FILE_NAME)
        ws_cars = wb["Cars"]

        price = None
        car_row = None

        for row in range(2, ws_cars.max_row + 1):

            if str(ws_cars.cell(row, 1).value) == car_id:

                price = float(
                    ws_cars.cell(row, 6).value
                )

                car_row = row
                break

        if price is None:
            messagebox.showerror(
                "Error",
                "Car ID not found"
            )
            return

        availability = ws_cars.cell(
            car_row,
            7
        ).value

        if availability != "Available":

            messagebox.showerror(
                "Error",
                "Car is not available"
            )
            return

        total_amount = price * total_days

        ws_rentals = wb["Rentals"]

        ws_rentals.append([
            entries["Rental ID"].get(),
            entries["Customer Name"].get(),
            entries["Contact"].get(),
            car_id,
            entries["Rental Date"].get(),
            entries["Return Date"].get(),
            total_days,
            total_amount
        ])

        ws_cars.cell(
            car_row,
            7
        ).value = "Rented"

        wb.save(FILE_NAME)

        messagebox.showinfo(
            "Rental Successful",
            f"Total Days: {total_days}\n"
            f"Total Amount: ₹{total_amount}"
        )

        window.destroy()

    tk.Button(
        window,
        text="Rent Car",
        width=20,
        command=rent_car
    ).pack(pady=25)


# ================= LOGOUT =================

def logout():

    root.destroy()

    global login_window
    login_window = tk.Tk()

    create_login_window()

    login_window.mainloop()


# ================= LOGIN GUI =================

def create_login_window():

    global username_entry
    global password_entry

    login_window.title(
        "Car Rental Management System - Login"
    )

    login_window.geometry("400x350")
    login_window.resizable(False, False)

    tk.Label(
        login_window,
        text="CAR RENTAL SYSTEM",
        font=("Arial", 20, "bold")
    ).pack(pady=30)

    tk.Label(
        login_window,
        text="Username"
    ).pack()

    username_entry = tk.Entry(
        login_window,
        width=30
    )
    username_entry.pack(pady=5)

    tk.Label(
        login_window,
        text="Password"
    ).pack()

    password_entry = tk.Entry(
        login_window,
        width=30,
        show="*"
    )
    password_entry.pack(pady=5)

    tk.Button(
        login_window,
        text="Login",
        width=20,
        command=login
    ).pack(pady=25)


# ================= START PROGRAM =================

create_excel()

login_window = tk.Tk()

create_login_window()

login_window.mainloop()